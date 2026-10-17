#!/usr/bin/env python3
"""
tools/excel_tests_umwandeln.py

Wandelt Eignungstests aus Excel (Vorlage „Test Durchführung / Verhalten /
Ergebnis“) in eine JSON-Datei für admin_tests_import.php um. Läuft lokal,
nicht auf dem Server. Braucht openpyxl (pip install openpyxl).

Aufruf:
    python tools/excel_tests_umwandeln.py <ausgabe.json> <datei.xlsx> [weitere.xlsx ...]

Erkennung (die Spalten verschieben sich je Datei, daher keine festen
Positionen):
- Testblatt: das erste Blatt, das nicht „Auswertung“ heißt und eine
  Kopfzeile mit „Verhalten“ hat.
- Name: Titel aus dem Blatt „Auswertung“ (A1), sonst erste Zeile des
  Testblatts. Beschreibung: Zelle A1 des Testblatts.
- Zielgruppe: „welpe“, wenn Dateiname oder Titel „Welpe“ enthält, sonst
  „erwachsen“.
- Aufgabe: Zeile mit einer Zahl in Spalte A. Titel ist die erste Textzelle
  links der Verhaltensspalte. Anleitung ist der weitere Text links der
  Verhaltensspalte bis zur nächsten Aufgabe.
- Abschnitt: Zeile mit Text in Spalte A ohne Zahl (z. B. „Wesen und
  Bindung“), außer Kopf- und Bemerkungszeilen.
- Ergebnisoptionen: Text in der Verhaltensspalte mit EXC, OK oder IND in
  derselben Zeile. Hat eine Aufgabe keine Optionen (Teil „Wesen und
  Bindung“), entstehen die Optionen „Exzellent (EXC)“ und
  „Indiskutabel (IND)“.
- Bewertung: EXC und OK = bestanden, IND = nicht bestanden. Die Abkürzung
  bleibt im Text der Option stehen.

Unklares landet in „warnungen“, die Datei wird trotzdem geschrieben.
"""

import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl fehlt: pip install openpyxl')

KATEGORIEN = {'EXC': 'bestanden', 'OK': 'bestanden', 'IND': 'nicht_bestanden'}
OHNE_OPTIONEN = [('Exzellent (EXC)', 'bestanden'), ('Indiskutabel (IND)', 'nicht_bestanden')]
KEINE_TITEL = {'EXC', 'OK', 'IND', 'STRESS?', 'ÜSH?', 'ÜS?', 'ZUSATZ', 'ERGEBNIS', 'VERHALTEN'}
MAX_TITEL = 200
MAX_OPTION = 200
MAX_ABSCHNITT = 150
MAX_ABSCHNITT_ERKANNT = 120  # längerer Text in Spalte A ist ein Hinweis, kein Abschnitt
MAX_KLAMMER_IM_TITEL = 40    # längere Klammerzusätze wandern vom Titel in die Anleitung
UNTERAUFGABE = re.compile(r'^AUFGABE\s+\d+\w*\s*:?\s*(.*)$', re.IGNORECASE)


def text(wert):
    """Zelle als bereinigter Text: Zeilenumbrüche zu Leerzeichen, doppelte Leerzeichen weg."""
    if wert is None:
        return ''
    return re.sub(r'\s+', ' ', str(wert)).strip()


def mehrzeilig(wert):
    """Wie text(), aber Zeilenumbrüche bleiben (für Beschreibungen)."""
    if wert is None:
        return ''
    zeilen = [re.sub(r'[ \t]+', ' ', z).strip() for z in str(wert).splitlines()]
    return '\n'.join(z for z in zeilen if z)


def titel_und_zusatz(titel):
    """Trennt lange Klammerzusätze vom Titel, z. B. „Sozialisierung 4 (gilt nicht für …)“."""
    zusaetze = [m.group(0) for m in re.finditer(r'\([^()]*\)', titel) if len(m.group(0)) > MAX_KLAMMER_IM_TITEL]
    for z in zusaetze:
        titel = titel.replace(z, ' ')
    return re.sub(r'\s+', ' ', titel).strip(), [z.strip('() ') for z in zusaetze]


def ist_nummer(wert):
    if isinstance(wert, (int, float)) and not isinstance(wert, bool):
        return float(wert).is_integer()
    return bool(re.fullmatch(r'\d+', text(wert)))


def kategorie(wert):
    k = text(wert).upper()
    return k if k in KATEGORIEN else None


def testblatt(wb):
    for ws in wb.worksheets:
        if ws.title.strip().lower() == 'auswertung':
            continue
        for r in range(1, min(ws.max_row, 10) + 1):
            for c in range(1, ws.max_column + 1):
                if text(ws.cell(r, c).value).lower() == 'verhalten':
                    return ws, r, c
    return None, None, None


def kategoriespalte(ws, ab_zeile, verhalten):
    """Spalte rechts der Verhaltensspalte, in der die meisten EXC/OK/IND stehen."""
    zaehler = {}
    for r in range(ab_zeile, ws.max_row + 1):
        if not text(ws.cell(r, verhalten).value):
            continue
        for c in range(verhalten + 1, min(verhalten + 4, ws.max_column) + 1):
            if kategorie(ws.cell(r, c).value):
                zaehler[c] = zaehler.get(c, 0) + 1
    return max(zaehler, key=zaehler.get) if zaehler else None


def umwandeln(pfad):
    warnungen = []
    datei = os.path.basename(pfad)
    wb = openpyxl.load_workbook(pfad, data_only=True)
    ws, kopfzeile, verhalten = testblatt(wb)
    if ws is None:
        return None, [f'{datei}: kein Testblatt mit Spalte „Verhalten“ gefunden, übersprungen']

    beschreibung = mehrzeilig(ws.cell(1, 1).value)
    name = ''
    if 'Auswertung' in wb.sheetnames:
        name = text(wb['Auswertung'].cell(1, 1).value)
    if not name:
        name = text(beschreibung.split('\n')[0])
    welpe = 'welpe' in datei.lower() or 'welpe' in name.lower()

    katspalte = kategoriespalte(ws, kopfzeile + 1, verhalten)
    if katspalte is None:
        warnungen.append(f'{datei}: keine Spalte mit EXC/OK/IND gefunden')

    aufgaben = []
    abschnitt = None
    aktuell = None
    hinweise = []

    def abschliessen():
        if aktuell is None:
            return
        if not aktuell['ergebnisse']:
            aktuell['ergebnisse'] = [{'bezeichnung': b, 'kategorie': k} for b, k in OHNE_OPTIONEN]
        aktuell['beschreibung'] = '\n'.join(aktuell.pop('_anleitung')) or None
        aktuell.pop('_basis', None)
        aufgaben.append(aktuell)

    for r in range(kopfzeile + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        a_text = text(a)
        if a_text.lower().startswith('bemerkung'):
            break

        if a is not None and ist_nummer(a):
            abschliessen()
            titel = ''
            for c in range(2, verhalten):
                t = text(ws.cell(r, c).value)
                if t and t.upper() not in KEINE_TITEL:
                    titel = t
                    break
            if not titel:
                warnungen.append(f'{datei}: Aufgabe {a_text} (Zeile {r}) ohne Titel')
                titel = f'Aufgabe {a_text}'
            roh_titel = titel
            titel, zusaetze = titel_und_zusatz(titel)
            if len(titel) > MAX_TITEL:
                warnungen.append(f'{datei}: Titel von Aufgabe {a_text} gekürzt')
                titel = titel[:MAX_TITEL]
            aktuell = {'nr': int(float(a)) if isinstance(a, (int, float)) else int(a_text),
                       'abschnitt': abschnitt, 'titel': titel, '_basis': titel,
                       '_anleitung': list(zusaetze), 'ergebnisse': []}
            start_spalte = None
            for c in range(2, verhalten):
                if text(ws.cell(r, c).value) == roh_titel:
                    start_spalte = c
                    break
            # Anleitung kann in derselben Zeile rechts vom Titel stehen
            for c in range((start_spalte or 1) + 1, verhalten):
                t = mehrzeilig(ws.cell(r, c).value)
                if t and text(t).upper() not in KEINE_TITEL:
                    aktuell['_anleitung'].append(t)
        elif a_text:
            if a_text.lower().startswith('test'):
                continue
            if len(a_text) > MAX_ABSCHNITT_ERKANNT:
                # Langer Text in Spalte A ist ein allgemeiner Hinweis zum Test
                hinweise.append(mehrzeilig(a))
                continue
            # Abschnittszeile, z. B. „Wesen und Bindung“
            abschliessen()
            aktuell = None
            abschnitt = a_text[:MAX_ABSCHNITT]
            continue
        elif aktuell is not None:
            for c in range(2, verhalten):
                t = mehrzeilig(ws.cell(r, c).value)
                if t and text(t).upper() not in KEINE_TITEL and text(t) != aktuell['_basis']:
                    aktuell['_anleitung'].append(t)

        if aktuell is None:
            continue
        option = text(ws.cell(r, verhalten).value)
        kat = kategorie(ws.cell(r, katspalte).value) if katspalte else None
        if option and option.lower() != 'verhalten':
            if kat is None:
                unter = UNTERAUFGABE.match(option)
                if unter:
                    # „AUFGABE 36a: nach 2 Minuten“: eigene Aufgabe mit eigenen Optionen
                    zusatz = unter.group(1).strip() or option
                    if aktuell['ergebnisse']:
                        eltern = aktuell
                        anleitung = list(eltern['_anleitung'])
                        basis = eltern['_basis']
                        abschliessen()
                        aktuell = {'nr': eltern['nr'], 'abschnitt': eltern['abschnitt'], '_basis': basis,
                                   'titel': f"{basis} – {zusatz}"[:MAX_TITEL],
                                   '_anleitung': anleitung, 'ergebnisse': []}
                    else:
                        aktuell['titel'] = f"{aktuell['_basis']} – {zusatz}"[:MAX_TITEL]
                else:
                    # Hinweis ohne Bewertung, z. B. „(Einige Hunde gehen zum Tester …)“
                    aktuell['_anleitung'].append('Hinweis: ' + mehrzeilig(ws.cell(r, verhalten).value))
                continue
            bezeichnung = f'{option} ({kat})'
            if len(bezeichnung) > MAX_OPTION:
                warnungen.append(f'{datei}: Aufgabe {aktuell["nr"]}, Zeile {r}: Option gekürzt')
                bezeichnung = option[:MAX_OPTION - len(kat) - 4] + f'… ({kat})'
            aktuell['ergebnisse'].append({'bezeichnung': bezeichnung, 'kategorie': KATEGORIEN[kat]})

    abschliessen()

    if not aufgaben:
        warnungen.append(f'{datei}: keine Aufgaben erkannt')
    titel = [a['titel'] for a in aufgaben]
    doppelt = sorted({t for t in titel if titel.count(t) > 1})
    if doppelt:
        warnungen.append(f'{datei}: Aufgabentitel mehrfach: {doppelt}')

    if hinweise:
        beschreibung = '\n'.join(([beschreibung] if beschreibung else []) + hinweise)
    test = {
        'name': name[:150],
        'beschreibung': beschreibung or None,
        'zielgruppe': 'welpe' if welpe else 'erwachsen',
        'quelle': datei,
        'aufgaben': [{k: v for k, v in a.items() if v is not None or k != 'abschnitt'} for a in aufgaben],
    }
    return test, warnungen


def main(argv):
    if len(argv) < 3:
        sys.exit(__doc__)
    ausgabe, dateien = argv[1], argv[2:]
    tests, warnungen = [], []
    for pfad in dateien:
        test, w = umwandeln(pfad)
        warnungen.extend(w)
        if test is not None:
            tests.append(test)
    daten = {'format': 'hunderassen-tests', 'version': 1, 'tests': tests, 'warnungen': warnungen}
    with open(ausgabe, 'w', encoding='utf-8') as f:
        json.dump(daten, f, ensure_ascii=False, indent=1)
    for t in tests:
        abschnitte = sorted({a.get('abschnitt') for a in t['aufgaben'] if a.get('abschnitt')})
        optionen = sum(len(a['ergebnisse']) for a in t['aufgaben'])
        print(f'{t["quelle"]}: "{t["name"]}" ({t["zielgruppe"]}), {len(t["aufgaben"])} Aufgaben, '
              f'{optionen} Optionen, Abschnitte: {abschnitte}')
    for w in warnungen:
        print('WARNUNG', w)
    print(f'{len(tests)} Tests nach {ausgabe}')


if __name__ == '__main__':
    main(sys.argv)
